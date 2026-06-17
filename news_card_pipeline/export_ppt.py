from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .utils import ensure_dir, truncate_text


def export_ppt(card_news_items: list[dict[str, Any]], output_path: Path) -> None:
    ensure_dir(output_path.parent)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for item in card_news_items:
        for card in item.get("cards", []):
            slide = prs.slides.add_slide(prs.slide_layouts[6])

            title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(11.9), Inches(0.9))
            title_frame = title_box.text_frame
            title_frame.text = card.get("title", "")
            title_run = title_frame.paragraphs[0].runs[0]
            title_run.font.bold = True
            title_run.font.size = Pt(30)

            body_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.5))
            body_frame = body_box.text_frame
            body_frame.word_wrap = True
            body_frame.text = truncate_text(card.get("body", ""), 420)
            body_paragraph = body_frame.paragraphs[0]
            body_paragraph.alignment = PP_ALIGN.LEFT
            body_paragraph.runs[0].font.size = Pt(22)

            footer = f"{item.get('source', '')} | {item.get('url', '')}"
            footer_box = slide.shapes.add_textbox(Inches(0.7), Inches(6.65), Inches(11.9), Inches(0.45))
            footer_frame = footer_box.text_frame
            footer_frame.text = truncate_text(footer, 180)
            footer_frame.paragraphs[0].runs[0].font.size = Pt(10)

    prs.save(output_path)
